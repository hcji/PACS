"""把两个规模化批次导出成 `result_1_robustness.xlsx` 那样的逐被试表。

产物（`results/data/`）：
- `result_2_univariate.xlsx`    单维度批：31 primitive × 2 程度 × 3 重复 + 通用退化对照 = 189 行
- `result_3_multivariable.xlsx` 组合批：C(31,2) = 465 对 × 3 重复 = 1395 行

## 格式从哪来

不是自己定的，是**照抄** `result_1_robustness.xlsx`：工作表名 `by_subject`、表头在第 1 行、
16 列（顺序与列名一致）、`row_id` = `<条件>#<被试序号 3 位>_r<重复>`、冻结 A2、
表头 宋体 11 加粗 + 顶对齐自动换行、第 1 行行高 57.6、各列宽照原样（A/B/C/D/F/G/H/I/J/K/P，
其余用默认 9.0）。

**刻意照抄的两处细节**（知道它们不理想，但一致性优先——这三个文件多半会被同一段代码读）：

1. `completeness` 与 7 个维度读数存成**文本**（原文件里是 `'0.3'` / `'0.45'`）。数字更好用
   （Excel 的「以文本形式存储的数字」告警、画图、求和都要先转换），但三个文件一起处理时
   类型不一致反而更容易出错，所以先跟原文件对齐。
2. `replicate` 与 `turns_run` 存成**整数**（原文件就是整数）。

## 行的来源

只导出**这两次实验自己跑出来的被试**：`result_2` 是 `scale_single` 的 189 个，
`result_3` 是 `scale_pair` 的 1395 个。完整心智基线（5 份，来自 Fig.1 主运行）与组合批的
可加性参照（`scale_single` 的 0.8 档）**不混进来** —— 它们属于别的批次，塞进同一个文件会让
`row_id` 前缀与出处对不上（§一致性优先于方便）。

除 `by_subject` 外各加一张聚合表：单维度批的剂量-反应/特异性判定、组合批的逐对交互效应。
它们不破坏原有的逐被试格式（第一张表仍然一模一样），但少了它们，`result_3` 就装不下它
这次实验真正要回答的东西。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_ROOT = Path(__file__).resolve().parents[2]
PACS = _ROOT / "pacs"
for _path in (str(PACS / "src"), str(PACS / "tools")):
    if _path not in sys.path:
        sys.path.insert(0, _path)

from pacs.phenotype import SCORED_CHANNELS, channel_label  # noqa: E402

SINGLE_RUN = PACS / "data" / "runs" / "scale_single"
PAIR_RUN = PACS / "data" / "runs" / "scale_pair"
OUT_DIR = _ROOT / "results" / "data"

SHEET = "by_subject"
DIMS = list(SCORED_CHANNELS)

#: 照抄 `result_1_robustness.xlsx` 的表头与版式
HEADERS = [
    "row_id", "mind_setting", "completeness", "replicate", "turns_run",
    "brain_model", "examiner_model", "analyst_model",
    *[f"{dim} [{channel_label(dim)}]" for dim in DIMS],
    "analyst_summary",
]
COLUMN_WIDTHS = {
    "A": 29.5555555555556, "B": 31.6666666666667, "C": 14.0, "D": 11.0,
    "F": 22.5555555555556, "G": 23.5555555555556, "H": 22.4444444444444,
    "I": 26.0, "J": 25.0, "K": 26.0, "P": 46.0,
}
HEADER_ROW_HEIGHT = 57.6
FONT_NAME = "宋体"


# ---------------------------------------------------------------------------
# 读取与字段提取
# ---------------------------------------------------------------------------


def load_records(run_dir: Path) -> list[dict]:
    """读一个批次的全部被试记录，按（条件, 被试序号, 重复）排序 —— 与 result_1 的顺序一致。"""
    records: list[dict] = []
    for path in sorted((run_dir / "cells").glob("*/subjects/*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["_path"] = str(path)
        records.append(payload)
    records.sort(key=lambda r: (
        str(r["condition"]["condition_id"]),
        int(r.get("subject_index") or 0),
        int((r["condition"].get("replicate") or 0)),
    ))
    return records


def row_id(record: dict) -> str:
    """`<条件>#<被试序号 3 位>_r<重复>`（照抄 result_1 的命名）。"""
    return (f"{record['condition']['condition_id']}"
            f"#{int(record.get('subject_index') or 0):03d}"
            f"_r{int(record['condition'].get('replicate') or 0)}")


def mind_setting(record: dict) -> str:
    """被削弱的过程与完整度。

    单因素：`belief_updating 0.80`；组合：`belief_updating + inhibition 0.80`
    （两者同剂量，剂量写在末尾，与 result_1 的「机制 + 程度」写法一致）；
    通用退化对照没有 primitive 被削弱，写它的参数名 `generic_degradation 0.50`。
    """
    params = record.get("parameters") or {}
    active = params.get("active") or {}
    if active:
        parts = " + ".join(sorted(active))
        level = float(next(iter(active.values())))
        return f"{parts} {level:.2f}"
    generic = float(params.get("generic_degradation") or 0.0)
    if generic > 0:
        return f"generic_degradation {generic:.2f}"
    return "intact (no perturbation)"


def completeness(record: dict) -> str:
    """程度值。照抄 result_1 的做法存成**文本**（原文件里是 `'0.3'`）。"""
    params = record.get("parameters") or {}
    active = params.get("active") or {}
    if active:
        return str(round(float(next(iter(active.values()))), 4))
    generic = float(params.get("generic_degradation") or 0.0)
    return str(round(generic, 4)) if generic > 0 else "1"


def subject_row(record: dict) -> list:
    condition = record.get("condition") or {}
    session = (record.get("dialogue") or {}).get("session") or {}
    analysis = record.get("analysis") or {}
    readings = analysis.get("readings") or {}
    return [
        row_id(record),
        mind_setting(record),
        completeness(record),
        int(condition.get("replicate") or 0),
        int(session.get("n_turns_run") or 0),
        condition.get("brain_model", ""),
        condition.get("examiner_model", ""),
        condition.get("analyst_model", ""),
        *["" if readings.get(dim) is None else str(round(float(readings[dim]), 4)) for dim in DIMS],
        analysis.get("conclusion") or "",
    ]


# ---------------------------------------------------------------------------
# 写表
# ---------------------------------------------------------------------------


def style_header(ws, headers: list) -> None:
    """表头：宋体 11 加粗、顶对齐 + 自动换行、行高 57.6 —— 逐项照抄 result_1。"""
    for column, text in enumerate(headers, start=1):
        cell = ws.cell(1, column, text)
        cell.font = Font(name=FONT_NAME, size=11, bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
    ws.freeze_panes = "A2"


def write_rows(ws, headers: list, rows: list, widths: dict) -> None:
    style_header(ws, headers)
    for row in rows:
        ws.append(row)
    # 数据区字体与 result_1 一致（宋体 11、不加粗、默认对齐）
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=11)
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        if letter in widths:
            ws.column_dimensions[letter].width = widths[letter]


def new_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)                     # 全部由写表函数创建，避免留一个空 Sheet
    wb.properties.creator = "Z.ai"
    return wb


def sheet_by_subject(wb, records: list) -> None:
    ws = wb.create_sheet(SHEET)
    write_rows(ws, HEADERS, [subject_row(r) for r in records], COLUMN_WIDTHS)


def _avg(values: list) -> object:
    """等权平均；缺值的通道不参与（不让 None 变成 0 而把均值拉低）。"""
    present = [float(v) for v in values if isinstance(v, (int, float))]
    return round(sum(present) / len(present), 4) if present else ""


def sheet_dose_and_specificity(wb, single_run: Path) -> int:
    """单维度批的聚合：逐（primitive × 程度）的数值 + 剂量-反应与特异性判定。

    `mean_reading` / `delta_mean` 是 7 个计分通道的等权平均 —— 与逐通道列同源，
    读者不必自己再横着算一遍。`dose_response_*` 是**该 primitive 跨两个程度**的拟合，
    所以同一 primitive 的两行会给出同一组拟合值（判定本身不是一个程度一个）。
    """
    result = json.loads((single_run / "analysis" / "scale_analysis.json").read_text(encoding="utf-8"))
    fits = {f["primitive_id"]: f for f in result["dose_response"]}
    spec: dict[tuple[str, float], dict] = {}
    for rows_at_dose in result["specificity"].values():
        for row in rows_at_dose:
            spec[(row["primitive_id"], float(row["dose"]))] = row

    headers = [
        "primitive_id", "primitive_zh", "dose", "n_replicates", "mean_reading", "delta_mean",
        *[f"{dim}_mean" for dim in DIMS],
        *[f"{dim}_delta" for dim in DIMS],
        "dominant_dim", "dominant_delta", "control_value", "control_delta",
        "specificity_verdict", "dose_response_verdict", "dose_response_slope",
        "dose_response_r2", "specificity_reason",
    ]
    rows = []
    for cell in result["cells"]:
        pid, dose = cell["primitive_id"], float(cell["dose"])
        channels = cell["channels"]
        verdict = spec.get((pid, dose), {})
        fit = fits.get(pid, {})
        means = [(channels.get(dim) or {}).get("mean") for dim in DIMS]
        deltas = [(channels.get(dim) or {}).get("delta") for dim in DIMS]
        rows.append([
            pid, cell["primitive_zh"], dose, cell["n_replicates"],
            _avg(means), _avg(deltas),
            *means, *deltas,
            verdict.get("dominant_dim", ""), verdict.get("dominant_delta", ""),
            verdict.get("control_value", ""), verdict.get("control_delta", ""),
            verdict.get("verdict", ""), fit.get("verdict", ""),
            fit.get("slope", ""), fit.get("r2", ""),
            verdict.get("reason", ""),
        ])
    rows.sort(key=lambda r: (r[2], r[0]))       # 先程度后名称，两个档各自成块

    widths = {"A": 34, "B": 16, "C": 8, "D": 12, "E": 14, "F": 12}
    start = len(widths) + 1
    for offset in range(2 * len(DIMS)):
        widths[get_column_letter(start + offset)] = 14
    tail = start + 2 * len(DIMS)
    for offset, width in enumerate([32, 14, 14, 12, 22, 20, 16, 10, 80]):
        widths[get_column_letter(tail + offset)] = width

    ws = wb.create_sheet("summary_by_cell")
    write_rows(ws, headers, rows, widths)
    return len(rows)


def sheet_interaction(wb, pair_run: Path) -> int:
    """组合批的聚合：逐对交互效应（这次实验真正要回答的东西）。

    `interaction_*` = 实测组合均值 − 两次单因子均值 + 基线均值（可加性下的期望）。
    正值 = 次加（两者一起反而没那么差），负值 = 超加。
    """
    result = json.loads((pair_run / "analysis" / "scale_pair_analysis.json").read_text(encoding="utf-8"))
    headers = [
        "primitive_a", "primitive_b", "label", "n_channels",
        "combined_mean", "baseline_mean", "delta_mean", "expected_additive_mean",
        "interaction_mean", "interaction_phenotype_mean",
        "n_superadditive", "n_additive", "n_subadditive",
        "most_interacting_channel", "most_interacting_kind",
        *[f"interaction_{dim}" for dim in DIMS],
    ]
    rows = []
    for pair in result["interactions"]["pairs"]:
        kinds = pair.get("kinds") or {}
        channels = pair.get("channels") or {}
        rows.append([
            pair["primitive_a"], pair["primitive_b"], pair["label"], pair["n_channels"],
            pair["combined_mean"], pair["baseline_mean"], pair["delta_mean"],
            pair["expected_additive_mean"],
            pair["interaction_mean"], pair["interaction_phenotype_mean"],
            kinds.get("superadditive", 0), kinds.get("additive", 0), kinds.get("subadditive", 0),
            pair["most_interacting_channel"], pair["most_interacting_kind"],
            *[(channels.get(dim) or {}).get("interaction", "") for dim in DIMS],
        ])
    rows.sort(key=lambda r: r[8])          # 交互量升序：最超加的在最前
    widths = {"A": 32, "B": 32, "C": 34, "D": 10, "E": 13, "F": 13, "G": 12, "H": 18,
              "I": 13, "J": 18, "K": 14, "L": 11, "M": 12, "N": 32, "O": 18}
    for offset in range(len(DIMS)):
        widths[get_column_letter(16 + offset)] = 20
    ws = wb.create_sheet("interaction_by_pair")
    write_rows(ws, headers, rows, widths)
    return len(rows)


# ---------------------------------------------------------------------------


def build_single() -> dict:
    records = load_records(SINGLE_RUN)
    wb = new_workbook()
    sheet_by_subject(wb, records)
    n_agg = sheet_dose_and_specificity(wb, SINGLE_RUN)
    path = OUT_DIR / "result_2_univariate.xlsx"
    wb.save(path)
    return {"file": str(path), "subjects": len(records), "agg_rows": n_agg,
            "conditions": len({r["condition"]["condition_id"] for r in records}),
            "settings": len({mind_setting(r) for r in records})}


def build_pair() -> dict:
    records = load_records(PAIR_RUN)
    wb = new_workbook()
    sheet_by_subject(wb, records)
    n_agg = sheet_interaction(wb, PAIR_RUN)
    path = OUT_DIR / "result_3_multivariable.xlsx"
    wb.save(path)
    return {"file": str(path), "subjects": len(records), "agg_rows": n_agg,
            "conditions": len({r["condition"]["condition_id"] for r in records}),
            "settings": len({mind_setting(r) for r in records})}


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, builder in (("result_2_univariate", build_single),
                          ("result_3_multivariable", build_pair)):
        info = builder()
        print(f"✓ {name}: {info['subjects']} 行被试、{info['conditions']} 个条件、"
              f"{info['settings']} 种设置、聚合表 {info['agg_rows']} 行")
        print(f"   {info['file']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
